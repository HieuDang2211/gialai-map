"""Nhập file sua_vi_tri_dia_diem.xlsx đã điền vào data/corrections/locations.json (mục "manual").

Chạy:  python -X utf8 api/scripts/apply-location-fixes.py <file_da_dien.xlsx> data/corrections/locations.json [--dry-run]
Sau đó: docker compose exec api node scripts/seed.js  (+ làm mới tile, xem README)

Chỉ nhận dòng có "Kết quả kiểm tra" và cột "Kiểm tra" = OK (tính lại tại đây, không tin giá trị đã lưu trong file).
  Vị trí đúng                    -> status verified
  Sửa vị trí                     -> status corrected + toạ độ mới
  Địa điểm không tồn tại         -> status hidden (ẩn khỏi bản đồ)
  Đã chuyển đi (chưa rõ nơi mới) -> status suspect
Mục manual đã có của địa điểm khác được giữ nguyên; trang /review ghi vào cùng mục này.
"""
import json
import math
import os
import sys
from datetime import date, datetime

import openpyxl

ACTION = {'Vị trí đúng': 'verified', 'Sửa vị trí': 'corrected', 'Địa điểm không tồn tại': 'hidden',
          'Đã chuyển đi (chưa rõ nơi mới)': 'suspect'}


def check(res, lat, lng, method, old_lat, old_lng):
    if res == 'Sửa vị trí':
        if not isinstance(lat, (int, float)) or not isinstance(lng, (int, float)):
            return 'Thiếu vĩ độ / kinh độ mới'
        if not (12.5 <= lat <= 15 and 107 <= lng <= 109.6):
            return 'Toạ độ ngoài Gia Lai'
        if not method:
            return 'Chưa chọn cách xác định'
        moved = math.hypot((lng - old_lng) * 111320 * math.cos(math.radians(old_lat)), (lat - old_lat) * 110540)
        if moved > 20000:
            return f'Dời {moved / 1000:.1f} km – kiểm tra lại'
        return None
    if lat not in (None, '') or lng not in (None, ''):
        return 'Có toạ độ mới nhưng kết quả không phải Sửa vị trí'
    return None


TRUSTED_TO_HIDE = {'Dữ liệu gialai-travel'}  # nguồn đã biết là kém tin cậy: được ẩn ngay theo báo cáo
BIG_MOVE_M = 1000


def hold_reason(e, src, moved, shared):
    """Chính sách --safe: thay đổi lớn trên dữ liệu chính thức / OSM chờ xác nhận thay vì áp ngay."""
    if e['status'] == 'hidden' and src not in TRUSTED_TO_HIDE:
        return 'báo không tồn tại nhưng nguồn ' + src + ' vẫn liệt kê'
    if e['status'] == 'corrected' and shared:
        return 'toạ độ mới trùng với địa điểm khác'
    if e['status'] == 'corrected' and moved > BIG_MOVE_M and src not in TRUSTED_TO_HIDE:
        return f'đề xuất dời {moved / 1000:.1f} km'
    return None


def write_held(path, held):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Chờ xác nhận'
    head = ['Mã hệ thống', 'Tên địa điểm', 'Nguồn dữ liệu', 'Phường / xã', 'Đề xuất', 'Lý do chờ', 'Vĩ độ đề xuất', 'Kinh độ đề xuất',
            'Xác nhận (Đồng ý / Không)', 'Ghi chú']
    ws.append(head)
    for h in held:
        ws.append([h['id'], h['name'], h['src'], h['ward'], h['proposal'], h['reason'], h.get('lat'), h.get('lng'), None, None])
    for i, w in enumerate([12, 34, 18, 18, 20, 36, 12, 12, 16, 26], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    dv = openpyxl.worksheet.datavalidation.DataValidation(type='list', formula1='"Đồng ý,Không"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'I2:I{len(held) + 1}')
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:J{len(held) + 1}'
    wb.save(path)


def main(xlsx, corrections, dry=False, safe=False):
    ws = openpyxl.load_workbook(xlsx, data_only=False)['Danh sách kiểm tra']
    hdr = [c.value for c in ws[4]]
    col = {h: i for i, h in enumerate(hdr)}
    entries, errors, stats, held = {}, [], {}, []
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[col['Mã hệ thống']]]
    new_points = {}
    for r in rows:
        if r[col['Kết quả kiểm tra']] == 'Sửa vị trí' and isinstance(r[col['Vĩ độ mới']], (int, float)):
            k = (round(r[col['Vĩ độ mới']], 5), round(r[col['Kinh độ mới']], 5))
            new_points[k] = new_points.get(k, 0) + 1
    for r in rows:
        res = r[col['Kết quả kiểm tra']]
        if not res:
            continue
        pid, name = r[col['Mã hệ thống']], r[col['Tên địa điểm']]
        if res not in ACTION:
            errors.append(f'{name}: kết quả lạ "{res}"')
            continue
        lat, lng = r[col['Vĩ độ mới']], r[col['Kinh độ mới']]
        method, when = r[col['Cách xác định']], r[col['Ngày kiểm tra']]
        problem = check(res, lat, lng, method, r[col['Vĩ độ hiện tại']], r[col['Kinh độ hiện tại']])
        if problem:
            errors.append(f'{name}: {problem}')
            continue
        day = when.strftime('%d/%m/%Y') if isinstance(when, (datetime, date)) else (str(when) if when else 'chưa ghi ngày')
        note = r[col['Ghi chú']]
        status = ACTION[res]
        e = {'status': status, 'name': name, 'reviewedAt': date.today().isoformat(), 'source': os.path.basename(xlsx)}
        if status == 'corrected':
            e.update(lng=round(float(lng), 7), lat=round(float(lat), 7), method=method,
                     note=f'Đã sửa vị trí ({method[:1].lower() + method[1:]}, {day})')
        elif status == 'verified':
            e['note'] = f'Đã kiểm tra vị trí {day}'
        elif status == 'hidden':
            e['note'] = f'Kiểm tra {day}: địa điểm không tồn tại'
        else:
            e['note'] = f'Kiểm tra {day}: đã chuyển đi, chưa rõ nơi mới'
        # ghi chú của người kiểm tra: chỉ lưu nội bộ, không hiển thị trên trang chi tiết
        if note:
            e['reviewNote'] = str(note)
        if safe:
            src = r[col['Nguồn dữ liệu']]
            moved = 0
            shared = False
            if status == 'corrected':
                olat, olng = r[col['Vĩ độ hiện tại']], r[col['Kinh độ hiện tại']]
                moved = math.hypot((e['lng'] - olng) * 111320 * math.cos(math.radians(olat)), (e['lat'] - olat) * 110540)
                shared = new_points.get((round(e['lat'], 5), round(e['lng'], 5)), 0) > 1
            why = hold_reason(e, src, moved, shared)
            if why:
                proposal = 'Ẩn (không tồn tại)' if status == 'hidden' else f'Dời {round(moved)} m'
                held.append({'id': pid, 'name': name, 'src': src, 'ward': r[col['Phường / xã']], 'proposal': proposal,
                             'reason': why, 'lat': e.get('lat'), 'lng': e.get('lng')})
                # chưa áp: bản đồ giữ nguyên, đề xuất nằm trong file chờ xác nhận
                stats['held'] = stats.get('held', 0) + 1
                continue
        entries[pid] = e
        stats[status] = stats.get(status, 0) + 1

    print(f'Nhận {len(entries)} dòng: {stats}; bỏ qua {len(errors)} dòng lỗi')
    for e in errors[:30]:
        print('  -', e)
    if held:
        from collections import Counter
        print('Chờ xác nhận theo lý do:', Counter(h['reason'].split(' nhưng')[0].split(' (')[0] if 'dời' not in h['reason'] else 'dời > 1 km' for h in held))
        if not dry:
            out = os.path.splitext(xlsx)[0] + '_cho_xac_nhan.xlsx'
            write_held(out, held)
            print(f'Danh sách chờ xác nhận: {out} ({len(held)} dòng)')
    if dry or not entries:
        return
    data = json.load(open(corrections, encoding='utf-8'))
    data['manual'] = {**data.get('manual', {}), **entries}
    tmp = corrections + '.tmp'
    json.dump(data, open(tmp, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    os.replace(tmp, corrections)
    print(f'Đã ghi {len(entries)} mục vào {corrections} (manual: {len(data["manual"])})')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2], '--dry-run' in sys.argv, '--safe' in sys.argv)
