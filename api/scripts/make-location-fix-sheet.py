"""Xuất Excel để kiểm tra / sửa vị trí các địa điểm (trừ quán ăn), rồi nhập lại bằng apply-location-fixes.py.

Chạy:  python -X utf8 api/scripts/make-location-fix-sheet.py <places.json> data/templates/sua_vi_tri_dia_diem.xlsx
  places.json: [{id, name, cat, slug, source, brand, address, ward, lat, lng, status, note}] (xuất từ bảng locations)
"""
import json
import sys
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = 'Arial'
TEAL = '0F766E'
HEAD = PatternFill('solid', fgColor=TEAL)
FIELD_HEAD = PatternFill('solid', fgColor='B45309')
FIELD_FILL = PatternFill('solid', fgColor='FFFBEB')
THIN = Side(style='thin', color='94A3B8')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F = Font(name=FONT, size=10)

SOURCE_LABEL = {'gialai-travel': 'Dữ liệu gialai-travel', 'osm': 'OpenStreetMap', 'vietcombank': 'Website Vietcombank',
                'agribank': 'Website Agribank', 'mb': 'Trang bản đồ MB', 'survey': 'Khảo sát'}
STATUS_LABEL = {'verified': 'Đã đối chiếu', 'corrected': 'Đã hiệu chỉnh tự động', 'unverified': 'Chưa đối chiếu',
                'suspect': 'Nghi sai'}
RESULTS = ['Vị trí đúng', 'Sửa vị trí', 'Địa điểm không tồn tại', 'Đã chuyển đi (chưa rõ nơi mới)']
METHODS = ['Đo GPS tại chỗ', 'Chọn trên ảnh vệ tinh', 'Khác']
PRIO_ORDER = {'Cao': 0, 'Trung bình': 1, 'Thấp': 2}


def priority(p):
    if p['status'] in ('suspect', 'unverified'):
        return 'Cao'
    if p['source'] == 'osm' or p['status'] == 'corrected':
        return 'Trung bình'
    return 'Thấp'


def build(places, out, title='KIỂM TRA VÀ SỬA VỊ TRÍ ĐỊA ĐIỂM – BẢN ĐỒ SỐ GIA LAI (trừ quán ăn)'):
    for p in places:
        p['prio'] = priority(p)
    places.sort(key=lambda p: (PRIO_ORDER[p['prio']], {'suspect': 0, 'unverified': 1}.get(p['status'], 2), p['cat'], p['ward'] or '', p['name']))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Danh sách kiểm tra'
    ws['A1'] = title
    ws['A1'].font = Font(name=FONT, size=14, bold=True, color=TEAL)
    ws['A2'] = (f'{len(places)} địa điểm, sắp theo Ưu tiên (Cao trước). Cột XANH: thông tin hiện tại – không sửa. '
                f'Cột CAM: điền kết quả kiểm tra. Xuất ngày {date.today():%d/%m/%Y}.')
    ws['A2'].font = Font(name=FONT, size=9, italic=True, color='475569')

    cols = [('STT', 5), ('Mã hệ thống', 12), ('Ưu tiên', 9), ('Tên địa điểm', 32), ('Danh mục', 14), ('Nguồn dữ liệu', 16),
            ('Địa chỉ hiện tại', 32), ('Phường / xã', 16), ('Vĩ độ hiện tại', 11), ('Kinh độ hiện tại', 11),
            ('Mức tin cậy hiện tại', 13), ('Ghi chú của hệ thống', 30), ('Xem vị trí hiện tại', 10),
            ('Kết quả kiểm tra', 20), ('Vĩ độ mới', 11), ('Kinh độ mới', 11), ('Cách xác định', 17), ('Ngày kiểm tra', 11),
            ('Người kiểm tra', 14), ('Ghi chú', 26), ('Dời bao xa (m)', 10), ('Kiểm tra', 30)]
    FIELD_FROM, FIELD_TO = 14, 20  # cột người dùng điền
    for j, (t, w) in enumerate(cols, start=1):
        c = ws.cell(row=4, column=j, value=t)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = FIELD_HEAD if FIELD_FROM <= j <= FIELD_TO else HEAD if j < FIELD_FROM else PatternFill('solid', fgColor='334155')
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[4].height = 34

    first = 5
    last = first + len(places) - 1
    for i, p in enumerate(places):
        r = first + i
        link = f"https://www.openstreetmap.org/?mlat={p['lat']:.6f}&mlon={p['lng']:.6f}#map=18/{p['lat']:.6f}/{p['lng']:.6f}"
        vals = [i + 1, p['id'], p['prio'], p['name'], p['cat'], SOURCE_LABEL.get(p['source'], p['source']), p['address'],
                p['ward'], float(p['lat']), float(p['lng']), STATUS_LABEL.get(p['status'], p['status']), p['note'], 'Bản đồ']
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F
            c.border = BORDER
            c.alignment = Alignment(wrap_text=j in (4, 7, 12), vertical='top')
        ws.cell(row=r, column=13).hyperlink = link
        ws.cell(row=r, column=13).font = Font(name=FONT, size=10, color='1D4ED8', underline='single')
        ws.cell(row=r, column=9).number_format = ws.cell(row=r, column=10).number_format = '0.000000'
        for j in range(FIELD_FROM, FIELD_TO + 1):
            c = ws.cell(row=r, column=j)
            c.fill = FIELD_FILL
            c.border = BORDER
            c.font = F
        ws.cell(row=r, column=15).number_format = ws.cell(row=r, column=16).number_format = '0.000000'
        ws.cell(row=r, column=18).number_format = 'dd/mm/yyyy'
        # Dời bao xa + kiểm tra (tự động)
        ws.cell(row=r, column=21, value=f'=IF(AND(ISNUMBER(O{r}),ISNUMBER(P{r})),ROUND(SQRT(((P{r}-J{r})*111320*COS(RADIANS(I{r})))^2+((O{r}-I{r})*110540)^2),0),"")')
        ws.cell(row=r, column=21).number_format = '#,##0'
        ws.cell(row=r, column=22, value=(
            f'=IF(N{r}="","",'
            f'IF(N{r}="Sửa vị trí",'
            f'IF(OR(O{r}="",P{r}=""),"Thiếu vĩ độ / kinh độ mới",'
            f'IF(OR(O{r}<12.5,O{r}>15,P{r}<107,P{r}>109.6),"Toạ độ ngoài Gia Lai (đảo vĩ/kinh độ?)",'
            f'IF(Q{r}="","Chọn cách xác định",'
            f'IF(U{r}>20000,"Dời hơn 20 km – kiểm tra lại","OK")))),'
            f'IF(AND(N{r}<>"Sửa vị trí",OR(O{r}<>"",P{r}<>"")),"Có toạ độ mới nhưng kết quả không phải Sửa vị trí","OK")))'))
        for j in (21, 22):
            ws.cell(row=r, column=j).font = Font(name=FONT, size=10, bold=(j == 22))
            ws.cell(row=r, column=j).border = BORDER

    dvs = [
        (DataValidation(type='list', formula1='"' + ','.join(RESULTS) + '"', allow_blank=True, showErrorMessage=True,
                        error='Chọn kết quả trong danh sách.'), f'N{first}:N{last}'),
        (DataValidation(type='decimal', operator='between', formula1='12.5', formula2='15', showErrorMessage=True,
                        error='Vĩ độ Gia Lai: 12,5 – 15,0 (vd 13.981234).'), f'O{first}:O{last}'),
        (DataValidation(type='decimal', operator='between', formula1='107', formula2='109.6', showErrorMessage=True,
                        error='Kinh độ Gia Lai: 107,0 – 109,6 (vd 108.004567).'), f'P{first}:P{last}'),
        (DataValidation(type='list', formula1='"' + ','.join(METHODS) + '"', allow_blank=True), f'Q{first}:Q{last}'),
        (DataValidation(type='date', operator='between', formula1='DATE(2025,1,1)', formula2='DATE(2030,12,31)',
                        showErrorMessage=True, error='Nhập ngày dd/mm/yyyy.'), f'R{first}:R{last}'),
    ]
    for dv, rng in dvs:
        ws.add_data_validation(dv)
        dv.add(rng)
    ws.conditional_formatting.add(f'V{first}:V{last}', FormulaRule(
        formula=[f'V{first}="OK"'], fill=PatternFill('solid', fgColor='DCFCE7'), font=Font(name=FONT, color='15803D', bold=True)))
    ws.conditional_formatting.add(f'V{first}:V{last}', FormulaRule(
        formula=[f'AND(V{first}<>"",V{first}<>"OK")'], fill=PatternFill('solid', fgColor='FEE2E2'), font=Font(name=FONT, color='B91C1C', bold=True)))
    ws.conditional_formatting.add(f'C{first}:C{last}', FormulaRule(
        formula=[f'C{first}="Cao"'], fill=PatternFill('solid', fgColor='FEE2E2'), font=Font(name=FONT, color='B91C1C', bold=True)))
    ws.freeze_panes = 'E5'
    ws.auto_filter.ref = f'A4:V{last}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '4:4'

    # ------------------------------------------------------------------ Tổng hợp
    ts = wb.create_sheet('Tổng hợp')
    ts['A1'] = 'TIẾN ĐỘ KIỂM TRA THEO DANH MỤC'
    ts['A1'].font = Font(name=FONT, size=13, bold=True, color=TEAL)
    head = ['Danh mục', 'Tổng', 'Ưu tiên Cao', 'Đã kiểm tra', 'Vị trí đúng', 'Sửa vị trí', 'Không tồn tại / chuyển đi', 'Dòng lỗi', '% đã kiểm tra']
    for j, h in enumerate(head, start=1):
        c = ts.cell(row=3, column=j, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = HEAD
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    for j, w in enumerate([22, 8, 10, 11, 11, 10, 16, 9, 12], start=1):
        ts.column_dimensions[get_column_letter(j)].width = w
    L = "'Danh sách kiểm tra'"
    cat_r, pr_r, res_r, chk_r = (f'{L}!${c}${first}:${c}${last}' for c in ('E', 'C', 'N', 'V'))
    cats = [c for c, _ in Counter(p['cat'] for p in places).most_common()]
    for i, cat in enumerate(cats, start=4):
        row = [cat, f'=COUNTIF({cat_r},A{i})', f'=COUNTIFS({cat_r},A{i},{pr_r},"Cao")', f'=COUNTIFS({cat_r},A{i},{res_r},"<>")',
               f'=COUNTIFS({cat_r},A{i},{res_r},"Vị trí đúng")', f'=COUNTIFS({cat_r},A{i},{res_r},"Sửa vị trí")',
               f'=COUNTIFS({cat_r},A{i},{res_r},"Địa điểm không tồn tại")+COUNTIFS({cat_r},A{i},{res_r},"Đã chuyển đi (chưa rõ nơi mới)")',
               f'=COUNTIFS({cat_r},A{i},{chk_r},"<>OK",{res_r},"<>")', f'=IF(B{i}=0,0,D{i}/B{i})']
        for j, v in enumerate(row, start=1):
            c = ts.cell(row=i, column=j, value=v)
            c.font = F
            c.border = BORDER
        ts.cell(row=i, column=9).number_format = '0%'
    tr = 4 + len(cats)
    ts.cell(row=tr, column=1, value='Tổng').font = Font(name=FONT, size=10, bold=True)
    for j in range(2, 9):
        Lc = get_column_letter(j)
        ts.cell(row=tr, column=j, value=f'=SUM({Lc}4:{Lc}{tr - 1})').font = Font(name=FONT, size=10, bold=True)
    ts.cell(row=tr, column=9, value=f'=IF(B{tr}=0,0,D{tr}/B{tr})').number_format = '0%'
    for j in range(1, 10):
        ts.cell(row=tr, column=j).border = BORDER

    # ------------------------------------------------------------------ Hướng dẫn
    hs = wb.create_sheet('Hướng dẫn', 0)
    hs.column_dimensions['A'].width = 115
    text = [
        ('HƯỚNG DẪN KIỂM TRA VÀ SỬA VỊ TRÍ', 't'),
        ('', ''),
        ('1. Làm theo thứ tự Ưu tiên: Cao (nghi sai / chưa đối chiếu) → Trung bình (OpenStreetMap, hiệu chỉnh tự động) → Thấp (website ngân hàng, đã đối chiếu).', ''),
        ('2. Bấm "Bản đồ" để xem vị trí hiện tại. Đến tận nơi hoặc đối chiếu trên ảnh vệ tinh.', ''),
        ('3. Chọn "Kết quả kiểm tra":', ''),
        ('    • Vị trí đúng – ghim đang đúng chỗ (để trống Vĩ độ / Kinh độ mới).', ''),
        ('    • Sửa vị trí – điền Vĩ độ mới, Kinh độ mới (≥ 6 chữ số thập phân, vd 13.981234 / 108.004567) và Cách xác định.', ''),
        ('    • Địa điểm không tồn tại – sẽ bị ẩn khỏi bản đồ.', ''),
        ('    • Đã chuyển đi – sẽ được đánh dấu "nghi sai" cho tới khi có địa chỉ mới.', ''),
        ('4. Cột "Kiểm tra" tự báo lỗi (thiếu toạ độ, toạ độ ngoài Gia Lai, đảo vĩ/kinh độ, dời quá xa). Chỉ dòng "OK" được nhập vào hệ thống.', ''),
        ('5. Cột "Dời bao xa (m)" tự tính khoảng cách từ vị trí cũ tới vị trí mới để tự soát lại.', ''),
        ('', ''),
        ('LẤY TOẠ ĐỘ ĐÚNG CÁCH', 'h'),
        ('• Đo GPS tại chỗ: đứng ở cổng/lối vào, chờ điện thoại báo độ chính xác ≤ 10 m rồi mới ghi; ghi đủ số thập phân điện thoại hiển thị, không làm tròn.', ''),
        ('• Chọn trên ảnh vệ tinh: phóng to tới khi thấy rõ mái nhà / cổng, đặt điểm ở lối vào chính.', ''),
        ('• Không chép toạ độ từ Google Maps hay ứng dụng khác; không tự cộng/trừ để "làm khác đi".', ''),
        ('• Ngân hàng / ATM / cây xăng: đặt điểm ở vị trí máy ATM, quầy giao dịch hoặc trụ bơm.', ''),
        ('', ''),
        ('SAU KHI ĐIỀN XONG', 'h'),
        ('Gửi lại file; hệ thống nhập bằng api/scripts/apply-location-fixes.py (chỉ các dòng "OK"), ghi vào data/corrections/locations.json mục "manual".', ''),
        ('Không xoá hoặc sửa cột "Mã hệ thống" – đó là khoá để nhập đúng địa điểm.', ''),
    ]
    for i, (t, k) in enumerate(text, start=1):
        c = hs.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, size=14 if k == 't' else 11 if k == 'h' else 10, bold=k in ('t', 'h'),
                      color=TEAL if k in ('t', 'h') else '000000')
        c.alignment = Alignment(wrap_text=True)

    wb.active = 1
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print(f'{len(places)} địa điểm -> {out};', dict(Counter(p['prio'] for p in places)))


if __name__ == '__main__':
    # tham số thứ 3 (tuỳ chọn): tiêu đề, vd khi chỉ xuất một khu vực
    build(json.load(open(sys.argv[1], encoding='utf-8')), sys.argv[2], *sys.argv[3:4])
