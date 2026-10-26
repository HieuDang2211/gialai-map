"""Tạo sổ khảo sát thực địa quán ăn TP Pleiku cũ (gộp nhiều nguồn, bỏ trùng, sắp theo phường + đường).

Chạy:
  python -X utf8 api/scripts/make-survey-list.py <danh_sach.json> <he_thong.json> <osm_reference.geojsonseq> <wards.json> <out.xlsx>
    danh_sach.json : danh sách tham khảo (data/templates/danh_sach_quan_an_pleiku.json)
    he_thong.json  : quán ăn đang có trong hệ thống (xuất từ bảng locations)
    osm_reference  : data/osm/reference.geojsonseq
    wards.json     : [{code, name, merged, g(GeoJSON)}] (xuất từ admin_units)

Sổ chỉ dùng để ĐI KHẢO SÁT: các cột "tham khảo" không được nhập/công bố; thông tin chính thức ghi vào
data/templates/mau_thu_thap_quan_an.xlsx sau khi đến tận nơi.
Không mang sang: đánh giá sao, lượt đánh giá, trạng thái mở cửa, link Google Maps.
"""
import difflib
import json
import math
import re
import sys
import unicodedata
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

FONT = 'Arial'
TEAL = '0F766E'
HEAD = PatternFill('solid', fgColor=TEAL)
FIELD_HEAD = PatternFill('solid', fgColor='B45309')
FIELD_FILL = PatternFill('solid', fgColor='FFFBEB')
FLAG_FILL = PatternFill('solid', fgColor='FEF3C7')
THIN = Side(style='thin', color='94A3B8')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F = Font(name=FONT, size=10)

# Phạm vi TP Pleiku cũ trong đơn vị hành chính mới
AREA = [23575, 23563, 23584, 23586, 23602, 23590]  # 5 phường trung tâm + Xã Biển Hồ
OLD_WARDS = ['Diên Hồng', 'Đống Đa', 'Hoa Lư', 'Hội Phú', 'Hội Thương', 'Ia Kring', 'Phù Đổng', 'Tây Sơn', 'Thắng Lợi',
             'Thống Nhất', 'Trà Bá', 'Yên Đỗ', 'Yên Thế', 'Chi Lăng']
RESULTS = ['Đã thu thập đủ', 'Chủ quán hẹn lại', 'Từ chối cung cấp', 'Đã đóng cửa', 'Không tìm thấy', 'Đã chuyển địa chỉ',
           'Không phải quán ăn']
NOT_FOOD = re.compile(r'\b(chợ|siêu thị|karaoke|khách sạn|hotel|bar|pub|cocktail)\b', re.I)


def norm(s):
    s = unicodedata.normalize('NFD', str(s or '')).replace('đ', 'd').replace('Đ', 'D')
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn').lower().replace('\xa0', ' ')
    return ' '.join(''.join(c if c.isalnum() else ' ' for c in s).split())


def dist(a, b):
    return math.hypot((a[0] - b[0]) * 111320 * math.cos(math.radians(a[1])), (a[1] - b[1]) * 110540)


def inside(pt, poly):
    x, y = pt
    c = False
    for ring in poly:
        for i in range(len(ring)):
            x1, y1 = ring[i - 1]
            x2, y2 = ring[i]
            if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
                c = not c
    return c


def first_point(c):
    while isinstance(c[0], list):
        c = c[0]
    return c


NUM_TOKEN = r'\d[\w/\-]*'
ADMIN_PART = re.compile(r'^(phường|xã|tp\.?|thành phố|tỉnh|gia lai|việt nam|\d{5,6})\b', re.I)


def split_address(addr):
    """Tách số nhà và tên đường.
    '22-24 Nguyễn Văn Trỗi, …' -> ('22-24', 'Nguyễn Văn Trỗi');  '58 101, Phạm Văn Đồng' -> ('58 101', 'Phạm Văn Đồng')
    '178, Sư Vạn Hạnh' -> ('178', 'Sư Vạn Hạnh');  'Hẻm 38 Nguyễn Thị Định' -> ('Hẻm 38', 'Nguyễn Thị Định')"""
    parts = [p.strip() for p in str(addr or '').replace('\xa0', ' ').split(',') if p.strip()]
    parts = [p for p in parts if not ADMIN_PART.match(p)] or ['']
    # tên toà nhà/TTTM đứng trước: "Coop Mart Pleiku, 21 Cách Mạng Tháng Tám" -> lấy đoạn có số nhà
    if not re.match(r'^(?:số\s*)?\d', parts[0], re.I):
        withnum = next((p for p in parts[1:] if re.match(rf'^(?:số\s*)?{NUM_TOKEN}\s+[^\d]', p, re.I)), None)
        if withnum:
            parts = [withnum]
    first = parts[0]
    m = re.match(rf'^(?:số\s*)?((?:{NUM_TOKEN}\s*)+?)\s*(?:đường\s+)?([^\d].*)?$', first, re.I)
    if m and m.group(1):
        num, street = m.group(1).strip(), (m.group(2) or '').strip()
        if not street and len(parts) > 1:  # số nhà đứng riêng một đoạn: "178, Sư Vạn Hạnh"
            street = re.sub(r'^đường\s+', '', parts[1], flags=re.I)
        return num, street
    m = re.match(rf'^(hẻm|kiệt|ngõ|kiot|ki ốt)\s+({NUM_TOKEN})\s*(.*)$', first, re.I)
    if m:
        return f'{m.group(1).capitalize()} {m.group(2)}', (m.group(3) or (parts[1] if len(parts) > 1 else '')).strip()
    return '', re.sub(r'^đường\s+', '', first, flags=re.I)


def category_hint(*texts):
    t = norm(' '.join(str(x or '') for x in texts))
    if re.search(r'ca phe|cafe|coffee', t):
        return 'Quán cà phê'
    if re.search(r'tra sua|nuoc ep|sinh to|giai khat|che\b|kem\b', t):
        return 'Quán nước'
    return 'Nhà hàng'


def load(argv):
    ref_list = json.load(open(argv[1], encoding='utf-8'))
    system = json.load(open(argv[2], encoding='utf-8'))
    wards = json.load(open(argv[4], encoding='utf-8'))
    by_code = {w['code']: w for w in wards}

    def ward_of(lng, lat):
        for w in wards:
            polys = w['g']['coordinates'] if w['g']['type'] == 'MultiPolygon' else [w['g']['coordinates']]
            if any(inside((lng, lat), p) for p in polys):
                return w
        return None

    items = []
    k = list(ref_list[0].keys())
    for r in ref_list:
        items.append({'name': str(r[k[0]]).replace('\xa0', ' ').strip(), 'kind': r[k[1]], 'old_ward': r[k[2]],
                      'address': r[k[3]], 'lng': float(r[k[9]]), 'lat': float(r[k[8]]), 'src': 'Danh sách tham khảo'})
    for r in system:
        items.append({'name': r['name'], 'kind': r['cat'], 'old_ward': '', 'address': r['address'],
                      'lng': r['lng'], 'lat': r['lat'], 'src': 'Hệ thống'})
    for line in open(argv[3], encoding='utf-8'):
        f = json.loads(line.strip().lstrip('\x1e'))
        p = f['properties']
        if p.get('amenity') not in ('restaurant', 'fast_food', 'food_court', 'cafe') or not (p.get('name:vi') or p.get('name')):
            continue
        lng, lat = first_point(f['geometry']['coordinates'])
        addr = ' '.join(x for x in (p.get('addr:housenumber'), p.get('addr:street')) if x)
        items.append({'name': p.get('name:vi') or p.get('name'), 'kind': 'Quán cà phê' if p['amenity'] == 'cafe' else 'Nhà hàng',
                      'old_ward': '', 'address': addr, 'lng': lng, 'lat': lat, 'src': 'OpenStreetMap'})
    return items, ward_of, by_code


def merge(items, ward_of):
    kept, dropped = [], []
    for it in items:
        w = ward_of(it['lng'], it['lat'])
        if not w or w['code'] not in AREA:
            if it['src'] == 'Danh sách tham khảo':
                dropped.append((it, f"Toạ độ nằm ở {w['name'] if w else 'ngoài tỉnh'}, ngoài TP Pleiku cũ"))
            continue
        it['ward'] = w
        twin = next((k for k in kept if norm(k['name']) == norm(it['name']) and dist((k['lng'], k['lat']), (it['lng'], it['lat'])) < 300
                     or difflib.SequenceMatcher(None, norm(k['name']), norm(it['name'])).ratio() >= 0.85
                     and dist((k['lng'], k['lat']), (it['lng'], it['lat'])) < 120), None)
        if twin:
            if it['src'] not in twin['srcs']:
                twin['srcs'].append(it['src'])
            if not twin['address'] and it['address']:
                twin['address'] = it['address']
            continue
        it['srcs'] = [it['src']]
        kept.append(it)
    return kept, dropped


def flags(items):
    for it in items:
        f = []
        num, street = split_address(it['address'])
        it['num'], it['street'] = num, street or '(chưa rõ đường)'
        if it['old_ward']:
            old = norm(str(it['old_ward']).replace('Phường', ''))
            if old not in norm(it['ward']['merged'] or '') and old not in norm(it['ward']['name']):
                f.append(f"Địa chỉ ghi {it['old_ward']} nhưng vị trí thuộc {it['ward']['name']}")
        if not num:
            f.append('Địa chỉ chưa có số nhà')
        if NOT_FOOD.search(f"{it['name']} {it['kind']}"):
            f.append('Có thể không phải quán ăn')
        if len(it['name']) > 55 or re.search(r'\d+\s*k\b|giá', it['name'], re.I):
            f.append('Tên có thể kèm quảng cáo – ghi lại tên trên biển hiệu')
        near = [o['name'] for o in items if o is not it and dist((o['lng'], o['lat']), (it['lng'], it['lat'])) < 15]
        if near:
            f.append('Trùng vị trí với: ' + ', '.join(near[:2]))
        if 'Hệ thống' in it['srcs']:
            f.append('Đã có trên bản đồ – kiểm tra & bổ sung')
        it['flags'] = '; '.join(f)
        it['hint'] = category_hint(it['kind'], it['name'])
    order = {c: i for i, c in enumerate(AREA)}
    items.sort(key=lambda x: (order[x['ward']['code']], norm(x['street']),
                              int(re.match(r'\d+', x['num']).group()) if re.match(r'\d+', x['num']) else 99999, norm(x['name'])))
    for i, it in enumerate(items, start=1):
        it['code'] = f'PK-{i:03d}'
    return items


def header(ws, row, titles, field_from=None, height=34):
    for j, t in enumerate(titles, start=1):
        c = ws.cell(row=row, column=j, value=t)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = FIELD_HEAD if field_from and j >= field_from else HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = BORDER
    ws.row_dimensions[row].height = height


def print_setup(ws, landscape=True, title_rows=None):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.center.text = 'Trang &P / &N'


def build(items, dropped, out):
    wb = Workbook()
    n = len(items)

    # ------------------------------------------------------------------ Danh sách khảo sát
    ws = wb.active
    ws.title = 'Danh sách khảo sát'
    ws['A1'] = 'SỔ KHẢO SÁT QUÁN ĂN – TP PLEIKU CŨ'
    ws['A1'].font = Font(name=FONT, size=14, bold=True, color=TEAL)
    ws['A2'] = (f'{n} địa điểm, sắp theo phường mới → tên đường → số nhà. Cột XANH: tham khảo (chưa xác minh, không công bố). '
                f'Cột CAM: điền khi đến nơi. Lập ngày {date.today():%d/%m/%Y}.')
    ws['A2'].font = Font(name=FONT, size=9, italic=True, color='475569')
    cols = [('STT', 5), ('Mã KS', 8), ('Tên quán (theo danh sách)', 30), ('Danh mục gợi ý', 12), ('Địa chỉ tham khảo', 30),
            ('Đường', 18), ('Phường / xã mới', 16), ('Phường cũ', 12), ('Cần kiểm tra tại chỗ', 34), ('Nguồn danh sách', 20),
            ('Xem vị trí', 9),
            ('Ngày đến', 10), ('Kết quả', 17), ('Tên đúng trên biển hiệu', 24), ('Đã đo GPS', 8), ('Chủ quán ký phiếu đồng ý', 10),
            ('Người khảo sát', 13), ('Ghi chú', 24)]
    FIELD_FROM = 12
    header(ws, 4, [c[0] for c in cols], FIELD_FROM)
    for j, (_, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    first = 5
    for i, it in enumerate(items):
        r = first + i
        link = f"https://www.openstreetmap.org/?mlat={it['lat']:.6f}&mlon={it['lng']:.6f}#map=19/{it['lat']:.6f}/{it['lng']:.6f}"
        vals = [i + 1, it['code'], it['name'], it['hint'], it['address'], it['street'], it['ward']['name'],
                (it['old_ward'] or '').replace('Phường ', ''), it['flags'], ', '.join(it['srcs']), 'Bản đồ']
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F
            c.border = BORDER
            c.alignment = Alignment(wrap_text=j in (3, 5, 6, 9, 10), vertical='top')
        ws.cell(row=r, column=11).hyperlink = link
        ws.cell(row=r, column=11).font = Font(name=FONT, size=10, color='1D4ED8', underline='single')
        if it['flags']:
            ws.cell(row=r, column=9).fill = FLAG_FILL
        for j in range(FIELD_FROM, len(cols) + 1):
            c = ws.cell(row=r, column=j)
            c.fill = FIELD_FILL
            c.border = BORDER
            c.font = F
        ws.cell(row=r, column=FIELD_FROM).number_format = 'dd/mm/yyyy'
    last = first + n - 1
    res_col = get_column_letter(FIELD_FROM + 1)
    dv = DataValidation(type='list', formula1='"' + ','.join(RESULTS) + '"', allow_blank=True, showErrorMessage=True,
                        error='Chọn kết quả trong danh sách.')
    yn = DataValidation(type='list', formula1='"Có,Không"', allow_blank=True)
    dd = DataValidation(type='date', operator='between', formula1='DATE(2025,1,1)', formula2='DATE(2030,12,31)',
                        showErrorMessage=True, error='Nhập ngày dd/mm/yyyy.')
    for v in (dv, yn, dd):
        ws.add_data_validation(v)
    dv.add(f'{res_col}{first}:{res_col}{last}')
    yn.add(f'{get_column_letter(FIELD_FROM + 3)}{first}:{get_column_letter(FIELD_FROM + 4)}{last}')
    dd.add(f'{get_column_letter(FIELD_FROM)}{first}:{get_column_letter(FIELD_FROM)}{last}')
    ws.conditional_formatting.add(f'{res_col}{first}:{res_col}{last}', FormulaRule(
        formula=[f'{res_col}{first}="Đã thu thập đủ"'], fill=PatternFill('solid', fgColor='DCFCE7'), font=Font(name=FONT, color='15803D', bold=True)))
    ws.conditional_formatting.add(f'{res_col}{first}:{res_col}{last}', FormulaRule(
        formula=[f'OR({res_col}{first}="Đã đóng cửa",{res_col}{first}="Không tìm thấy",{res_col}{first}="Không phải quán ăn",{res_col}{first}="Từ chối cung cấp")'],
        fill=PatternFill('solid', fgColor='FEE2E2'), font=Font(name=FONT, color='B91C1C')))
    ws.freeze_panes = 'D5'
    ws.auto_filter.ref = f'A4:{get_column_letter(len(cols))}{last}'
    print_setup(ws, landscape=True, title_rows='4:4')

    # ------------------------------------------------------------------ Quán phát hiện thêm
    ws2 = wb.create_sheet('Quán phát hiện thêm')
    ws2['A1'] = 'QUÁN PHÁT HIỆN THÊM KHI ĐI KHẢO SÁT (chưa có trong danh sách)'
    ws2['A1'].font = Font(name=FONT, size=13, bold=True, color=TEAL)
    ws2['A2'] = ('Ưu tiên các phường cũ chưa có quán nào trong danh sách: Đống Đa, Yên Thế, Chi Lăng; và Thắng Lợi (mới có 11). '
                 'Mã KS bắt đầu từ PK-M001.')
    ws2['A2'].font = Font(name=FONT, size=9, italic=True, color='475569')
    cols2 = [('Mã KS', 9), ('Tên trên biển hiệu', 30), ('Danh mục', 13), ('Số nhà', 8), ('Đường', 20), ('Phường / xã mới', 18),
             ('Phường cũ', 12), ('Ngày đến', 10), ('Kết quả', 17), ('Đã đo GPS', 8), ('Chủ quán ký phiếu đồng ý', 10),
             ('Người khảo sát', 13), ('Ghi chú', 26)]
    header(ws2, 4, [c[0] for c in cols2], 8)
    for j, (_, w) in enumerate(cols2, start=1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    for r in range(5, 205):
        ws2.cell(row=r, column=1, value=f'PK-M{r - 4:03d}').font = F
        for j in range(1, len(cols2) + 1):
            c = ws2.cell(row=r, column=j)
            c.border = BORDER
            c.font = F
            if j >= 2:
                c.fill = FIELD_FILL
        ws2.cell(row=r, column=8).number_format = 'dd/mm/yyyy'
    for spec, rng in ((('"Nhà hàng,Quán cà phê,Quán nước"'), 'C5:C204'), ('"' + ','.join(f'Phường {w}' for w in ['Pleiku', 'Diên Hồng', 'Thống Nhất', 'Hội Phú', 'An Phú']) + ',Xã Biển Hồ"', 'F5:F204'),
                      ('"' + ','.join(OLD_WARDS) + '"', 'G5:G204'), ('"' + ','.join(RESULTS) + '"', 'I5:I204'), ('"Có,Không"', 'J5:K204')):
        v = DataValidation(type='list', formula1=spec, allow_blank=True)
        ws2.add_data_validation(v)
        v.add(rng)
    ws2.freeze_panes = 'C5'
    print_setup(ws2, landscape=True, title_rows='4:4')

    # ------------------------------------------------------------------ Tiến độ
    ws3 = wb.create_sheet('Tiến độ')
    ws3['A1'] = 'TIẾN ĐỘ KHẢO SÁT THEO PHƯỜNG / XÃ'
    ws3['A1'].font = Font(name=FONT, size=13, bold=True, color=TEAL)
    t = ['Phường / xã mới', 'Trong danh sách', 'Đã đến', 'Đã thu thập đủ', 'Hẹn lại', 'Từ chối', 'Đóng cửa / không thấy / không phải quán ăn',
         'Phát hiện thêm', '% hoàn thành']
    header(ws3, 3, t)
    for j, w in enumerate([22, 12, 10, 12, 10, 10, 18, 12, 12], start=1):
        ws3.column_dimensions[get_column_letter(j)].width = w
    L = "'Danh sách khảo sát'"
    wcol, rcol = f'{L}!$G${first}:$G${last}', f'{L}!${res_col}${first}:${res_col}${last}'
    names = [f'Phường {w}' for w in ['Pleiku', 'Diên Hồng', 'Thống Nhất', 'Hội Phú', 'An Phú']] + ['Xã Biển Hồ']
    for i, name in enumerate(names, start=4):
        cells = [name,
                 f'=COUNTIF({wcol},A{i})',
                 f'=COUNTIFS({wcol},A{i},{rcol},"<>")',
                 f'=COUNTIFS({wcol},A{i},{rcol},"Đã thu thập đủ")',
                 f'=COUNTIFS({wcol},A{i},{rcol},"Chủ quán hẹn lại")',
                 f'=COUNTIFS({wcol},A{i},{rcol},"Từ chối cung cấp")',
                 f'=COUNTIFS({wcol},A{i},{rcol},"Đã đóng cửa")+COUNTIFS({wcol},A{i},{rcol},"Không tìm thấy")+COUNTIFS({wcol},A{i},{rcol},"Không phải quán ăn")+COUNTIFS({wcol},A{i},{rcol},"Đã chuyển địa chỉ")',
                 f"=COUNTIF('Quán phát hiện thêm'!$F$5:$F$204,A{i})",
                 f'=IF(B{i}=0,0,C{i}/B{i})']
        for j, v in enumerate(cells, start=1):
            c = ws3.cell(row=i, column=j, value=v)
            c.font = F
            c.border = BORDER
        ws3.cell(row=i, column=9).number_format = '0%'
    tr = 4 + len(names)
    ws3.cell(row=tr, column=1, value='Tổng').font = Font(name=FONT, size=10, bold=True)
    for j in range(2, 9):
        L2 = get_column_letter(j)
        c = ws3.cell(row=tr, column=j, value=f'=SUM({L2}4:{L2}{tr - 1})')
        c.font = Font(name=FONT, size=10, bold=True)
        c.border = BORDER
    c = ws3.cell(row=tr, column=9, value=f'=IF(B{tr}=0,0,C{tr}/B{tr})')
    c.number_format = '0%'
    c.font = Font(name=FONT, size=10, bold=True)
    c.border = BORDER
    ws3.cell(row=tr, column=1).border = BORDER
    ws3.cell(row=tr + 2, column=1, value='Bảng tự cập nhật theo cột "Kết quả" của sheet Danh sách khảo sát và sheet Quán phát hiện thêm.').font = Font(name=FONT, size=9, italic=True, color='475569')

    # ------------------------------------------------------------------ Phiếu cung cấp thông tin (in A4)
    build_paper_form(wb.create_sheet('Phiếu cung cấp thông tin'))

    # ------------------------------------------------------------------ Hướng dẫn
    hs = wb.create_sheet('Hướng dẫn', 0)
    hs.column_dimensions['A'].width = 110
    lines = [
        ('HƯỚNG DẪN ĐI KHẢO SÁT QUÁN ĂN – TP PLEIKU CŨ', 'title'),
        ('', ''),
        ('CHUẨN BỊ', 'h'),
        ('1. Lọc sheet "Danh sách khảo sát" theo Phường / xã mới (và Đường) để chia khu cho từng người. Danh sách đã sắp theo đường, số nhà để đi một lượt.', ''),
        ('2. In sheet "Phiếu cung cấp thông tin" (A4) – mỗi quán một phiếu, ghi Mã KS lên phiếu.', ''),
        ('3. Mang điện thoại có GPS để đo vị trí ngay tại cửa quán.', ''),
        ('', ''),
        ('TẠI QUÁN', 'h'),
        ('4. Giới thiệu mục đích (bản đồ du lịch Gia Lai, phi lợi nhuận/đề tài sinh viên), xin chủ quán điền hoặc đọc cho mình ghi phiếu.', ''),
        ('5. Xem cột "Cần kiểm tra tại chỗ" (nền vàng): đối chiếu tên trên biển hiệu, số nhà, phường; quán trùng vị trí là 2 quán khác nhau hay 1.', ''),
        ('6. Đo GPS, ghi vĩ độ/kinh độ vào phiếu. Chỉ chụp ảnh khi chủ quán đồng ý.', ''),
        ('7. Xin chữ ký ở phần "Đồng ý công khai thông tin". Không ký = vẫn ghi nhận, nhưng KHÔNG công khai số điện thoại / ảnh.', ''),
        ('8. Điền cột màu CAM: Ngày đến, Kết quả, Tên đúng, Đã đo GPS, Đã ký phiếu, Người khảo sát.', ''),
        ('', ''),
        ('SAU KHI KHẢO SÁT', 'h'),
        ('9. Nhập thông tin từ phiếu vào file mau_thu_thap_quan_an.xlsx (một quán một dòng), ghi Mã KS vào cột "Ghi chú nội bộ".', ''),
        ('10. Quán mới gặp trên đường: ghi vào sheet "Quán phát hiện thêm" (mã PK-M…) và làm phiếu như bình thường.', ''),
        ('11. Sheet "Tiến độ" tự tổng hợp kết quả theo phường.', ''),
        ('', ''),
        ('LƯU Ý VỀ DỮ LIỆU', 'h'),
        ('• Các cột màu XANH (tên, địa chỉ, danh mục gợi ý, vị trí) chỉ để TÌM ĐẾN QUÁN: chưa xác minh, có thể sai, KHÔNG nhập vào hệ thống.', ''),
        ('• Một phần danh sách lấy từ Google Maps: không sao chép đánh giá, mô tả, ảnh hay giờ mở cửa của Google vào hệ thống.', ''),
        ('• Thông tin chính thức là thông tin ghi được tại quán và chủ quán đồng ý cung cấp.', ''),
        ('', ''),
        ('PHẠM VI', 'h'),
        (f'• {len(items)} địa điểm trong 5 phường trung tâm (Pleiku, Diên Hồng, Thống Nhất, Hội Phú, An Phú) và Xã Biển Hồ.', ''),
        ('• Nguồn danh sách: danh sách tham khảo bạn cung cấp, dữ liệu đang có trên bản đồ, OpenStreetMap – đã gộp các quán trùng.', ''),
        (f'• {len(dropped)} dòng bị loại vì vị trí ngoài TP Pleiku cũ – xem sheet "Đã loại".', ''),
        ('• Danh sách CHƯA đầy đủ: các phường cũ Đống Đa, Yên Thế, Chi Lăng gần như chưa có quán – cần đi quét thực địa.', ''),
    ]
    for i, (txt, kind) in enumerate(lines, start=1):
        c = hs.cell(row=i, column=1, value=txt)
        c.font = (Font(name=FONT, size=14, bold=True, color=TEAL) if kind == 'title' else
                  Font(name=FONT, size=11, bold=True, color=TEAL) if kind == 'h' else F)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    print_setup(hs, landscape=False)

    # ------------------------------------------------------------------ Đã loại
    ws5 = wb.create_sheet('Đã loại')
    header(ws5, 1, ['Tên', 'Địa chỉ ghi trong danh sách', 'Lý do loại'])
    for j, w in enumerate([34, 50, 50], start=1):
        ws5.column_dimensions[get_column_letter(j)].width = w
    for i, (it, why) in enumerate(dropped, start=2):
        for j, v in enumerate([it['name'], it['address'], why], start=1):
            c = ws5.cell(row=i, column=j, value=v)
            c.font = F
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical='top')

    wb.active = 1
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)


def build_paper_form(ws):
    """Phiếu A4 in ra giấy để chủ quán điền / người khảo sát ghi."""
    for col, w in zip('ABCDEFGH', [16, 12, 12, 12, 12, 12, 12, 14]):
        ws.column_dimensions[col].width = w
    box = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    r = 1

    def title(text, size=13):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, size=size, bold=True, color=TEAL)
        c.alignment = Alignment(horizontal='center' if size > 12 else 'left', vertical='center')
        ws.row_dimensions[r].height = 22
        r += 1

    def field(label, span_label=2, height=22):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span_label)
        ws.merge_cells(start_row=r, start_column=span_label + 1, end_row=r, end_column=8)
        c = ws.cell(row=r, column=1, value=label)
        c.font = Font(name=FONT, size=10, bold=True)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        for col in range(1, 9):
            ws.cell(row=r, column=col).border = box
        ws.row_dimensions[r].height = height
        r += 1

    def checks(label, options):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, size=10, bold=True)
        ws.cell(row=r, column=3, value='     '.join(f'☐ {o}' for o in options)).font = F
        for col in range(1, 9):
            ws.cell(row=r, column=col).border = box
        ws.cell(row=r, column=1).alignment = Alignment(vertical='center', wrap_text=True)
        ws.cell(row=r, column=3).alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[r].height = 22
        r += 1

    title('PHIẾU CUNG CẤP THÔNG TIN QUÁN ĂN – BẢN ĐỒ SỐ GIA LAI', 14)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.cell(row=r, column=1, value='Mã KS: ____________     Ngày: ___ / ___ / 20___     Người khảo sát: ______________________').font = F
    r += 2
    title('A. THÔNG TIN QUÁN', 11)
    field('Tên trên biển hiệu')
    checks('Danh mục', ['Nhà hàng', 'Quán cà phê', 'Quán nước'])
    field('Loại món chính / món đặc trưng', height=32)
    field('Số nhà – Tên đường')
    field('Phường / xã (mới)')
    field('Vĩ độ – Kinh độ (đo GPS tại cửa)', height=30)
    field('Số điện thoại quán')
    field('Website / Facebook')
    r += 1
    title('B. GIỜ MỞ CỬA (ghi HH:MM-HH:MM hoặc "Nghỉ")', 11)
    for j, d in enumerate(['', 'Thứ 2', 'Thứ 3', 'Thứ 4', 'Thứ 5', 'Thứ 6', 'Thứ 7', 'Chủ nhật'], start=1):
        c = ws.cell(row=r, column=j, value=d)
        c.font = Font(name=FONT, size=10, bold=True)
        c.alignment = Alignment(horizontal='center')
        c.border = box
    ws.cell(row=r + 1, column=1, value='Giờ mở – đóng').font = Font(name=FONT, size=10, bold=True)
    for j in range(1, 9):
        ws.cell(row=r + 1, column=j).border = box
    ws.row_dimensions[r + 1].height = 26
    r += 3
    title('C. GIÁ & TIỆN ÍCH', 11)
    field('Giá mỗi người (đồng): từ … đến …', height=30)
    checks('Tiện ích', ['Wifi', 'Máy lạnh', 'Đỗ ô tô', 'Thẻ / QR', 'Giao hàng', 'Có món chay'])
    field('Số chỗ ngồi (ước lượng)')
    checks('Giấy chứng nhận ATTP', ['Có', 'Không', 'Không rõ'])
    field('Giới thiệu ngắn về quán (≤ 160 ký tự)', height=44)
    checks('Ảnh', ['Người khảo sát chụp', 'Chủ quán gửi ảnh', 'Không có ảnh'])
    r += 1
    title('D. ĐỒNG Ý CÔNG KHAI THÔNG TIN', 11)
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 2, end_column=8)
    c = ws.cell(row=r, column=1, value=(
        'Tôi đồng ý để thông tin ở mục A, B, C (kể cả số điện thoại và ảnh nếu có) được hiển thị miễn phí trên Bản đồ số Gia Lai '
        '(đề tài sinh viên Đại học Duy Tân). Tôi có thể yêu cầu sửa hoặc gỡ thông tin bất kỳ lúc nào qua người khảo sát.'))
    c.font = F
    c.alignment = Alignment(wrap_text=True, vertical='top')
    for rr in range(r, r + 3):
        for col in range(1, 9):
            ws.cell(row=rr, column=col).border = box
    r += 3
    checks('Chủ quán', ['Đồng ý', 'Không đồng ý công khai SĐT/ảnh'])
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
    ws.cell(row=r, column=1, value='Họ tên chủ quán / người đại diện').font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=5, value='Chữ ký').font = Font(name=FONT, size=10, bold=True)
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 3, end_column=4)
    ws.merge_cells(start_row=r + 1, start_column=5, end_row=r + 3, end_column=8)
    for rr in range(r, r + 4):
        for col in range(1, 9):
            ws.cell(row=rr, column=col).border = box
    r += 4
    ws.cell(row=r + 1, column=1, value='Ghi chú của người khảo sát:').font = Font(name=FONT, size=10, bold=True)
    ws.merge_cells(start_row=r + 2, start_column=1, end_row=r + 4, end_column=8)
    for rr in range(r + 2, r + 5):
        for col in range(1, 9):
            ws.cell(row=rr, column=col).border = box
    ws.print_area = f'A1:H{r + 4}'
    print_setup(ws, landscape=False)
    ws.page_setup.fitToHeight = 1
    ws.sheet_view.showGridLines = False


if __name__ == '__main__':
    items, ward_of, _ = load(sys.argv)
    kept, dropped = merge(items, ward_of)
    kept = flags(kept)
    build(kept, dropped, sys.argv[5])
    from collections import Counter
    print(f'{len(kept)} địa điểm, loại {len(dropped)}; theo nguồn:',
          Counter(', '.join(k['srcs']) for k in kept).most_common(), '; theo phường:', Counter(k['ward']['name'] for k in kept).most_common(),
          '; có cần kiểm tra:', sum(1 for k in kept if k['flags']))
