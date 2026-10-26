"""Tạo file Excel mẫu để thu thập thông tin quán ăn / cà phê và nhập vào hệ thống.

Chạy:  python -I -X utf8 api/scripts/make-eatery-template.py data/templates/_wards.txt data/templates/mau_thu_thap_quan_an.xlsx
  _wards.txt: "mã|tên đầy đủ|sáp nhập từ" (xuất từ bảng admin_units)

Hàng 3 của sheet "Dữ liệu" là khoá cột (khớp trường trong bảng locations) để script nhập đọc tự động.
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

FONT = 'Arial'
TEAL = '0F766E'
GROUP_FILL = PatternFill('solid', fgColor='CCFBF1')
HEAD_FILL = PatternFill('solid', fgColor=TEAL)
REQ_FILL = PatternFill('solid', fgColor='B45309')  # tiêu đề cột bắt buộc
KEY_FILL = PatternFill('solid', fgColor='F1F5F9')
EXAMPLE_FILL = PatternFill('solid', fgColor='E2E8F0')
INPUT_REQ_FILL = PatternFill('solid', fgColor='FFFBEB')  # ô bắt buộc cần điền
THIN = Side(style='thin', color='CBD5E1')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FIRST_ROW = 5   # hàng nhập liệu đầu tiên (hàng 4 = ví dụ)
LAST_ROW = 504  # 500 dòng nhập

CATEGORIES = ['Nhà hàng', 'Quán cà phê', 'Quán nước']
CUISINES = ['Cơm', 'Phở / bún / mì / hủ tiếu', 'Lẩu / nướng', 'Hải sản', 'Đặc sản Tây Nguyên', 'Bánh mì / ăn sáng',
            'Ăn vặt', 'Chay', 'Gà / vịt', 'Dê / bò', 'Món Hoa / Hàn / Nhật', 'Đồ uống / giải khát', 'Khác']
PRICE = ['Bình dân', 'Tầm trung', 'Cao cấp']
YESNO = ['Có', 'Không']
YESNO3 = ['Có', 'Không', 'Không rõ']
CONSENT = ['Có', 'Không', 'Chưa hỏi']
METHOD = ['GPS tại chỗ', 'Chọn trên bản đồ hệ thống', 'OpenStreetMap', 'Khác']
SOURCE = ['Khảo sát thực địa', 'Chủ quán cung cấp', 'Cơ quan quản lý', 'OpenStreetMap']
PHOTO = ['Ảnh tự chụp', 'Chủ quán cho phép', 'Không có ảnh']
DAYS = [('mon', 'Thứ 2'), ('tue', 'Thứ 3'), ('wed', 'Thứ 4'), ('thu', 'Thứ 5'), ('fri', 'Thứ 6'), ('sat', 'Thứ 7'), ('sun', 'Chủ nhật')]

# (nhóm, khoá, tiêu đề, bắt buộc, độ rộng, kiểu, danh sách, ví dụ, ghi chú)
COLUMNS = [
    ('Thông tin cơ bản', 'name', 'Tên quán', True, 30, 'text', None, '(VÍ DỤ) Cơm gà Phố Núi',
     'Tên ghi trên biển hiệu. Không thêm địa chỉ vào tên.'),
    ('Thông tin cơ bản', 'category', 'Danh mục', True, 14, 'list', 'DS_DanhMuc', 'Nhà hàng',
     'Danh mục hiển thị trên bản đồ: Nhà hàng / Quán cà phê / Quán nước.'),
    ('Thông tin cơ bản', 'cuisine', 'Loại món chính', False, 20, 'list', 'DS_LoaiMon', 'Gà / vịt',
     'Chọn trong danh sách; dùng để lọc và tìm kiếm.'),
    ('Thông tin cơ bản', 'signature_dishes', 'Món đặc trưng', False, 26, 'text', None, 'Cơm gà xé, gỏi gà, canh chua',
     'Cách nhau bằng dấu phẩy. Người dùng tìm theo các món này.'),
    ('Vị trí', 'house_number', 'Số nhà', False, 9, 'text', None, '25',
     'Để trống nếu không có số nhà.'),
    ('Vị trí', 'street', 'Tên đường / thôn', True, 20, 'text', None, 'Trần Hưng Đạo',
     'Chỉ ghi tên đường, không ghi "đường".'),
    ('Vị trí', 'ward', 'Phường / xã (sau sáp nhập)', True, 24, 'list', 'DS_PhuongXa', 'Phường Pleiku',
     'Theo đơn vị hành chính MỚI (từ 7/2025). Sheet "Danh sách" có cột "Sáp nhập từ" để tra tên cũ.'),
    ('Vị trí', 'lat', 'Vĩ độ', True, 11, 'lat', None, 13.98123,
     'Số thập phân, ít nhất 5 chữ số sau dấu chấm. Gia Lai: 12,5 – 15,0.'),
    ('Vị trí', 'lng', 'Kinh độ', True, 11, 'lng', None, 108.00456,
     'Số thập phân. Gia Lai: 107,0 – 109,6.'),
    ('Vị trí', 'location_method', 'Cách lấy toạ độ', True, 18, 'list', 'DS_CachLayToaDo', 'GPS tại chỗ',
     'Ưu tiên đứng ở cửa quán đo bằng GPS điện thoại.'),
    ('Liên hệ', 'phone', 'Số điện thoại', False, 15, 'phone', None, '0905 123 456',
     'Số của quán (không ghi số cá nhân khi chủ quán chưa đồng ý).'),
    ('Liên hệ', 'website', 'Website / Facebook', False, 24, 'text', None, 'facebook.com/comgaphonui',
     'Đường dẫn trang chính thức của quán.'),
    ('Liên hệ', 'email', 'Email', False, 20, 'text', None, '',
     ''),
    *[('Giờ mở cửa', f'hours_{k}', v, False, 12, 'hours', None, ('06:00-21:30' if k != 'sun' else '07:00-21:00'),
       'Dạng HH:MM-HH:MM (vd 06:00-21:30), "24/24" hoặc "Nghỉ". Mở 2 ca: ghi ca dài nhất, ca còn lại ghi ở Ghi chú.')
      for k, v in DAYS],
    ('Giá', 'price_range', 'Mức giá', False, 11, 'list', 'DS_MucGia', 'Bình dân',
     'Bình dân: dưới 50.000đ/người · Tầm trung: 50.000–200.000đ · Cao cấp: trên 200.000đ.'),
    ('Giá', 'price_min', 'Giá thấp nhất (đ/người)', False, 13, 'money', None, 35000,
     'Số nguyên, đơn vị đồng.'),
    ('Giá', 'price_max', 'Giá cao nhất (đ/người)', False, 13, 'money', None, 70000,
     'Số nguyên, đơn vị đồng. Phải lớn hơn hoặc bằng giá thấp nhất.'),
    ('Tiện ích', 'wifi', 'Wifi', False, 7, 'list', 'DS_CoKhong', 'Có', ''),
    ('Tiện ích', 'air_con', 'Máy lạnh', False, 9, 'list', 'DS_CoKhong', 'Không', ''),
    ('Tiện ích', 'car_parking', 'Đỗ ô tô', False, 8, 'list', 'DS_CoKhong', 'Có', ''),
    ('Tiện ích', 'card_qr', 'Thẻ / QR', False, 8, 'list', 'DS_CoKhong', 'Có', 'Nhận thanh toán thẻ hoặc chuyển khoản QR.'),
    ('Tiện ích', 'delivery', 'Giao hàng', False, 9, 'list', 'DS_CoKhong', 'Có', ''),
    ('Tiện ích', 'vegetarian', 'Có món chay', False, 9, 'list', 'DS_CoKhong', 'Không', ''),
    ('Tiện ích', 'seats', 'Số chỗ ngồi', False, 9, 'int', None, 40, 'Ước lượng.'),
    ('Mô tả & hình ảnh', 'short_description', 'Mô tả ngắn (≤160 ký tự)', False, 34, 'short', None,
     'Quán cơm gà lâu năm gần Quảng trường Đại Đoàn Kết, phần ăn đầy đặn, giá bình dân.',
     'Hiện ngay dưới tên quán. Tối đa 160 ký tự.'),
    ('Mô tả & hình ảnh', 'description', 'Giới thiệu chi tiết', False, 36, 'text', None, '',
     'Không chép mô tả hay đánh giá từ trang khác (Google Maps, Foody…).'),
    ('Mô tả & hình ảnh', 'photos', 'Link ảnh (cách nhau bằng ;)', False, 30, 'text', None, '',
     'Chỉ ảnh tự chụp hoặc chủ quán cho phép. Không lấy ảnh từ Google Maps.'),
    ('Mô tả & hình ảnh', 'photo_permission', 'Quyền dùng ảnh', False, 15, 'list', 'DS_QuyenAnh', 'Ảnh tự chụp', ''),
    ('Nguồn & xác minh', 'source', 'Nguồn thông tin', True, 18, 'list', 'DS_Nguon', 'Khảo sát thực địa',
     'Nơi lấy thông tin. Không ghi "Google Maps": dữ liệu Google không được sao chép vào hệ thống.'),
    ('Nguồn & xác minh', 'collected_at', 'Ngày thu thập', True, 12, 'date', None, date(2026, 10, 17),
     'Ngày/tháng/năm.'),
    ('Nguồn & xác minh', 'collector', 'Người thu thập', True, 16, 'text', None, 'Nguyễn Văn A',
     'Họ tên người điền dòng này.'),
    ('Nguồn & xác minh', 'verified_onsite', 'Đã đến tận nơi', False, 10, 'list', 'DS_CoKhong', 'Có', ''),
    ('Nguồn & xác minh', 'food_safety_cert', 'Giấy chứng nhận ATTP', False, 12, 'list', 'DS_CoKhongRo', 'Không rõ',
     'Giấy chứng nhận cơ sở đủ điều kiện an toàn thực phẩm.'),
    ('Nguồn & xác minh', 'owner_consent', 'Chủ quán đồng ý công khai', False, 13, 'list', 'DS_DongY', 'Có',
     'Bắt buộc "Có" trước khi công khai số điện thoại / ảnh do chủ quán cung cấp.'),
    ('Nguồn & xác minh', 'notes', 'Ghi chú nội bộ', False, 26, 'text', None, 'Mở thêm ca tối 17:00-21:30',
     'Không hiển thị cho người dùng.'),
]

ARIAL = Font(name=FONT, size=10)


def build(wards_path, out_path):
    wards = [line.rstrip('\n').split('|') for line in open(wards_path, encoding='utf-8') if line.strip()]
    wb = Workbook()

    # ------------------------------------------------------------------ Danh sách (dropdown)
    ls = wb.active
    ls.title = 'Danh sách'
    lists = {
        'DS_DanhMuc': CATEGORIES, 'DS_LoaiMon': CUISINES, 'DS_MucGia': PRICE, 'DS_CoKhong': YESNO,
        'DS_CoKhongRo': YESNO3, 'DS_DongY': CONSENT, 'DS_CachLayToaDo': METHOD, 'DS_Nguon': SOURCE, 'DS_QuyenAnh': PHOTO,
    }
    col = 1
    for name, values in lists.items():
        c = ls.cell(row=1, column=col, value=name.replace('DS_', ''))
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = HEAD_FILL
        for i, v in enumerate(values, start=2):
            ls.cell(row=i, column=col, value=v).font = ARIAL
        letter = get_column_letter(col)
        wb.defined_names[name] = DefinedName(name, attr_text=f"'Danh sách'!${letter}$2:${letter}${len(values) + 1}")
        ls.column_dimensions[letter].width = max(14, max(len(v) for v in values) + 2)
        col += 1
    # phường/xã: mã + tên + sáp nhập từ
    col += 1
    for j, h in enumerate(['Mã', 'Phường / xã (sau sáp nhập)', 'Sáp nhập từ (tên cũ)']):
        c = ls.cell(row=1, column=col + j, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = HEAD_FILL
    for i, (code, name, merged) in enumerate(wards, start=2):
        ls.cell(row=i, column=col, value=int(code)).font = ARIAL
        ls.cell(row=i, column=col + 1, value=name).font = ARIAL
        ls.cell(row=i, column=col + 2, value=merged).font = ARIAL
    wl = get_column_letter(col + 1)
    wb.defined_names['DS_PhuongXa'] = DefinedName('DS_PhuongXa', attr_text=f"'Danh sách'!${wl}$2:${wl}${len(wards) + 1}")
    ls.column_dimensions[get_column_letter(col)].width = 8
    ls.column_dimensions[wl].width = 26
    ls.column_dimensions[get_column_letter(col + 2)].width = 60
    ls.freeze_panes = 'A2'

    # ------------------------------------------------------------------ Dữ liệu
    ws = wb.create_sheet('Dữ liệu', 0)
    ncol = len(COLUMNS) + 1  # + cột kiểm tra
    letters = {}
    for i, (group, key, title, req, width, kind, lst, example, note) in enumerate(COLUMNS, start=1):
        L = get_column_letter(i)
        letters[key] = L
        g = ws.cell(row=1, column=i, value=group)
        g.font = Font(name=FONT, size=10, bold=True, color=TEAL)
        g.fill = GROUP_FILL
        h = ws.cell(row=2, column=i, value=f'{title} *' if req else title)
        h.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        h.fill = REQ_FILL if req else HEAD_FILL
        h.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        h.border = BORDER
        if note:
            h.comment = Comment(note, 'Mẫu thu thập')
        k = ws.cell(row=3, column=i, value=key)
        k.font = Font(name=FONT, size=8, italic=True, color='64748B')
        k.fill = KEY_FILL
        e = ws.cell(row=4, column=i, value=example if example != '' else None)
        e.font = Font(name=FONT, size=10, italic=True, color='475569')
        e.fill = EXAMPLE_FILL
        e.border = BORDER
        ws.column_dimensions[L].width = width

        rng = f'{L}{FIRST_ROW}:{L}{LAST_ROW}'
        fmt = {'lat': '0.000000', 'lng': '0.000000', 'money': '#,##0', 'int': '0', 'date': 'dd/mm/yyyy', 'phone': '@'}.get(kind)
        for r in range(4, LAST_ROW + 1):
            cell = ws.cell(row=r, column=i)
            if r >= FIRST_ROW:
                cell.font = ARIAL
                cell.border = BORDER
                if req:
                    cell.fill = INPUT_REQ_FILL
            if fmt:
                cell.number_format = fmt
            if kind in ('text', 'short'):
                cell.alignment = Alignment(wrap_text=False, vertical='top')

        dv = None
        if kind == 'list':
            dv = DataValidation(type='list', formula1=f'={lst}', allow_blank=True, showErrorMessage=True,
                                errorTitle='Giá trị không hợp lệ', error='Chọn một giá trị trong danh sách.')
        elif kind == 'lat':
            dv = DataValidation(type='decimal', operator='between', formula1='12.5', formula2='15', showErrorMessage=True,
                                errorTitle='Vĩ độ ngoài Gia Lai', error='Vĩ độ phải trong khoảng 12,5 – 15,0 (vd 13.98123).')
        elif kind == 'lng':
            dv = DataValidation(type='decimal', operator='between', formula1='107', formula2='109.6', showErrorMessage=True,
                                errorTitle='Kinh độ ngoài Gia Lai', error='Kinh độ phải trong khoảng 107,0 – 109,6 (vd 108.00456).')
        elif kind == 'money':
            dv = DataValidation(type='whole', operator='between', formula1='0', formula2='100000000', showErrorMessage=True,
                                error='Nhập số nguyên (đồng), vd 35000.')
        elif kind == 'int':
            dv = DataValidation(type='whole', operator='between', formula1='0', formula2='10000', showErrorMessage=True,
                                error='Nhập số nguyên.')
        elif kind == 'date':
            dv = DataValidation(type='date', operator='between', formula1='DATE(2020,1,1)', formula2='DATE(2035,12,31)',
                                showErrorMessage=True, error='Nhập ngày dạng dd/mm/yyyy.')
        elif kind == 'short':
            dv = DataValidation(type='textLength', operator='lessThanOrEqual', formula1='160', showErrorMessage=True,
                                error='Mô tả ngắn tối đa 160 ký tự.')
        elif kind == 'hours':
            c0 = f'{L}{FIRST_ROW}'
            dv = DataValidation(
                type='custom',
                formula1=(f'OR({c0}="",{c0}="Nghỉ",{c0}="24/24",AND(LEN({c0})=11,MID({c0},3,1)=":",MID({c0},6,1)="-",'
                          f'MID({c0},9,1)=":",ISNUMBER(--LEFT({c0},2)),ISNUMBER(--MID({c0},4,2)),'
                          f'ISNUMBER(--MID({c0},7,2)),ISNUMBER(--RIGHT({c0},2))))'),
                showErrorMessage=True, errorTitle='Sai định dạng giờ',
                error='Nhập dạng HH:MM-HH:MM (vd 06:00-21:30), "24/24" hoặc "Nghỉ".')
        if dv:
            ws.add_data_validation(dv)
            dv.add(rng)

    # Cột kiểm tra tự động
    C = get_column_letter(ncol)
    ws.cell(row=1, column=ncol, value='Kiểm tra').font = Font(name=FONT, size=10, bold=True, color=TEAL)
    ws.cell(row=1, column=ncol).fill = GROUP_FILL
    hc = ws.cell(row=2, column=ncol, value='Kết quả kiểm tra (tự động)')
    hc.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
    hc.fill = PatternFill('solid', fgColor='334155')
    hc.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    hc.comment = Comment('Không sửa cột này. "OK" = dòng đủ điều kiện nhập vào hệ thống.', 'Mẫu thu thập')
    ws.cell(row=3, column=ncol, value='check').font = Font(name=FONT, size=8, italic=True, color='64748B')
    ws.cell(row=3, column=ncol).fill = KEY_FILL
    ws.column_dimensions[C].width = 34

    req_keys = [c[1] for c in COLUMNS if c[3]]
    labels = {c[1]: c[2] for c in COLUMNS}
    first, last = get_column_letter(1), get_column_letter(len(COLUMNS))
    for r in range(4, LAST_ROW + 1):
        L = letters
        missing = '&'.join(f'IF({L[k]}{r}="","{labels[k]}; ","")' for k in req_keys)
        formula = (
            f'=IF(COUNTA({first}{r}:{last}{r})=0,"",'
            f'IF(LEN({missing})>0,"Thiếu: "&{missing},'
            f'IF(OR({L["lat"]}{r}<12.5,{L["lat"]}{r}>15,{L["lng"]}{r}<107,{L["lng"]}{r}>109.6),"Toạ độ ngoài Gia Lai (đảo vĩ/kinh độ?)",'
            f'IF(AND(ISNUMBER({L["price_min"]}{r}),ISNUMBER({L["price_max"]}{r}),{L["price_max"]}{r}<{L["price_min"]}{r}),"Giá cao nhất nhỏ hơn giá thấp nhất",'
            f'IF(COUNTIF(${L["name"]}${FIRST_ROW}:${L["name"]}${LAST_ROW},{L["name"]}{r})>1,"Trùng tên với dòng khác – kiểm tra lại",'
            f'"OK")))))'
        )
        c = ws.cell(row=r, column=ncol, value=formula)
        c.font = Font(name=FONT, size=10, bold=True) if r >= FIRST_ROW else Font(name=FONT, size=10, italic=True, color='475569')
        c.border = BORDER
        if r == 4:
            c.fill = EXAMPLE_FILL
    ws.conditional_formatting.add(
        f'{C}{FIRST_ROW}:{C}{LAST_ROW}',
        FormulaRule(formula=[f'AND({C}{FIRST_ROW}<>"",{C}{FIRST_ROW}<>"OK")'], fill=PatternFill('solid', fgColor='FEE2E2'),
                    font=Font(name=FONT, color='B91C1C', bold=True)))
    ws.conditional_formatting.add(
        f'{C}{FIRST_ROW}:{C}{LAST_ROW}',
        FormulaRule(formula=[f'{C}{FIRST_ROW}="OK"'], fill=PatternFill('solid', fgColor='DCFCE7'),
                    font=Font(name=FONT, color='15803D', bold=True)))

    ws.row_dimensions[2].height = 42
    ws.freeze_panes = f'B{FIRST_ROW}'
    ws.auto_filter.ref = f'A2:{C}{LAST_ROW}'
    ws.cell(row=4, column=1).comment = Comment('Dòng VÍ DỤ (nền xám): xem cách điền rồi xoá hoặc bỏ qua. '
                                               'Hệ thống tự bỏ dòng có tên bắt đầu bằng "(VÍ DỤ)".', 'Mẫu thu thập')

    # ------------------------------------------------------------------ Hướng dẫn
    hs = wb.create_sheet('Hướng dẫn', 0)
    hs.column_dimensions['A'].width = 28
    hs.column_dimensions['B'].width = 11
    hs.column_dimensions['C'].width = 24
    hs.column_dimensions['D'].width = 30
    hs.column_dimensions['E'].width = 70
    rows = [
        ('MẪU THU THẬP THÔNG TIN QUÁN ĂN / CÀ PHÊ – BẢN ĐỒ SỐ GIA LAI', 'title'),
        ('Dùng để khảo sát và nhập quán ăn, quán cà phê, quán nước vào hệ thống. Mỗi quán là một dòng trong sheet "Dữ liệu".', 'p'),
        ('', 'p'),
        ('CÁCH ĐIỀN', 'h'),
        ('1. Điền từ hàng 5 của sheet "Dữ liệu". Hàng 4 (nền xám) là ví dụ; hàng 3 (chữ nhỏ) là khoá hệ thống – không sửa.', 'p'),
        ('2. Cột tiêu đề màu CAM có dấu * là BẮT BUỘC (ô nhập nền vàng nhạt). Cột tiêu đề màu XANH là tuỳ chọn – điền càng đủ càng tốt.', 'p'),
        ('3. Ô có mũi tên: chọn trong danh sách. Rê chuột lên tiêu đề cột để xem hướng dẫn chi tiết.', 'p'),
        ('4. Cột "Kết quả kiểm tra" tự báo dòng thiếu thông tin hoặc sai: chỉ dòng "OK" (xanh) mới được nhập vào hệ thống.', 'p'),
        ('5. Toạ độ: đứng ở cửa quán, mở bản đồ trên điện thoại, lấy vĩ độ/kinh độ (vd 13.98123, 108.00456). Vĩ độ ~13–14, kinh độ ~108.', 'p'),
        ('6. Phường/xã ghi theo đơn vị MỚI sau sáp nhập. Không nhớ tên mới: tra cột "Sáp nhập từ" trong sheet "Danh sách".', 'p'),
        ('', 'p'),
        ('LƯU Ý VỀ NGUỒN DỮ LIỆU', 'h'),
        ('• KHÔNG sao chép thông tin, mô tả, đánh giá hay ảnh từ Google Maps, Foody, Facebook của người khác: điều khoản của các nền tảng này cấm, và hệ thống không được phép công bố.', 'p'),
        ('• Nguồn hợp lệ: khảo sát thực địa, chủ quán tự cung cấp, danh sách của cơ quan quản lý, OpenStreetMap.', 'p'),
        ('• Số điện thoại và ảnh do chủ quán cung cấp chỉ được công khai khi cột "Chủ quán đồng ý công khai" = Có.', 'p'),
        ('', 'p'),
        ('DANH SÁCH CỘT', 'h'),
    ]
    r = 1
    for text, kind in rows:
        c = hs.cell(row=r, column=1, value=text)
        if kind == 'title':
            c.font = Font(name=FONT, size=14, bold=True, color=TEAL)
        elif kind == 'h':
            c.font = Font(name=FONT, size=11, bold=True, color=TEAL)
        else:
            c.font = ARIAL
        r += 1
    for j, h in enumerate(['Cột', 'Bắt buộc', 'Khoá hệ thống', 'Ví dụ', 'Hướng dẫn'], start=1):
        c = hs.cell(row=r, column=j, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = HEAD_FILL
        c.border = BORDER
    r += 1
    for group, key, title, req, width, kind, lst, example, note in COLUMNS:
        ex = example.strftime('%d/%m/%Y') if isinstance(example, date) else example
        vals = [f'{group} › {title}', 'Có' if req else '', key, ex, note]
        for j, v in enumerate(vals, start=1):
            c = hs.cell(row=r, column=j, value=v)
            c.font = Font(name=FONT, size=10, bold=(j == 2 and req), color=('B45309' if j == 2 and req else '000000'))
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1
    hs.freeze_panes = 'A2'

    wb.active = wb.index(ws)
    # openpyxl không lưu giá trị đã tính -> yêu cầu Excel/LibreOffice tính lại cột kiểm tra khi mở file
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(f'Đã tạo {out_path}: {len(COLUMNS)} cột, {LAST_ROW - FIRST_ROW + 1} dòng nhập, {len(wards)} phường/xã')


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
