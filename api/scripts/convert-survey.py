"""Chuyển sổ khảo sát đã điền (xlsx) -> data/raw/survey_quan_an_pleiku.json để seed.js nhập.

Chạy:  python -X utf8 api/scripts/convert-survey.py <so_da_khao_sat.xlsx> <so_goc.xlsx> data/raw/survey_quan_an_pleiku.json
  so_goc.xlsx: sổ khảo sát ban đầu (để lấy toạ độ tham khảo và so với toạ độ "GPS" đã điền).

Mỗi dòng ghi lại:
  gps_independent = False khi toạ độ ghi trong sổ trùng toạ độ tham khảo (< 0,5 m) -> chưa phải toạ độ đo tại chỗ
  shared_point    = True khi có quán khác cùng toạ độ (< 3 m)
"""
import json
import math
import re
import sys
from datetime import date, datetime

import openpyxl

RESULT_ACTION = {
    'Đã thu thập đủ': 'publish',
    'Chủ quán hẹn lại': 'publish',      # quán có thật, còn thiếu thông tin chi tiết
    'Từ chối cung cấp': 'publish',      # quán có thật; không công khai SĐT/ảnh
    'Đã đóng cửa': 'close',
    'Không tìm thấy': 'close',
    'Không phải quán ăn': 'close',
    'Đã chuyển địa chỉ': 'moved',
}


def dist(a, b):
    return math.hypot((a[0] - b[0]) * 111320 * math.cos(math.radians(a[1])), (a[1] - b[1]) * 110540)


def main(done_path, orig_path, out_path):
    done = openpyxl.load_workbook(done_path)['Danh sách khảo sát']
    orig = openpyxl.load_workbook(orig_path)['Danh sách khảo sát']
    ref = {}
    for i, r in enumerate(orig.iter_rows(min_row=5, values_only=True), start=5):
        m = re.search(r'mlat=([\d.]+)&mlon=([\d.]+)', orig.cell(row=i, column=11).hyperlink.target or '')
        ref[r[1]] = (float(m.group(2)), float(m.group(1))) if m else None

    hdr = [c.value for c in done[4]]
    col = {h: i for i, h in enumerate(hdr)}
    gps_col = next(h for h in hdr if h and ('GPS' in h))
    rows, problems = [], []
    for r in done.iter_rows(min_row=5, values_only=True):
        code = r[col['Mã KS']]
        if not code:
            continue
        g = re.match(r'\s*(-?[\d.]+)\s*[,;]\s*(-?[\d.]+)', str(r[col[gps_col]] or ''))
        if g:
            lat, lng = float(g.group(1)), float(g.group(2))
        elif ref.get(code):
            lng, lat = ref[code]
        else:
            problems.append(f'{code}: không có toạ độ')
            continue
        if not (12.5 <= lat <= 15 and 107 <= lng <= 109.6):
            problems.append(f'{code}: toạ độ ngoài Gia Lai ({lat}, {lng})')
            continue
        visited = r[col['Ngày đến']]
        result = (r[col['Kết quả']] or '').strip()
        rows.append({
            'code': code,
            'name': (r[col['Tên đúng trên biển hiệu']] or r[col['Tên quán (theo danh sách)']] or '').strip(),
            'listedName': (r[col['Tên quán (theo danh sách)']] or '').strip(),
            'category': r[col['Danh mục gợi ý']] or 'Nhà hàng',
            'address': r[col['Địa chỉ tham khảo']],
            'lng': lng,
            'lat': lat,
            'visitedAt': visited.date().isoformat() if isinstance(visited, datetime) else (visited.isoformat() if isinstance(visited, date) else visited),
            'result': result,
            'action': RESULT_ACTION.get(result, 'skip'),
            'consent': (r[col['Chủ quán ký phiếu đồng ý']] or '') == 'Có',
            'surveyor': r[col['Người khảo sát']],
            'note': r[col['Ghi chú']],
            'inSystem': 'Hệ thống' in str(r[col['Nguồn danh sách']] or ''),
            'gpsIndependent': not (ref.get(code) and dist(ref[code], (lng, lat)) < 0.5),
        })
    for a in rows:
        a['sharedPoint'] = any(b is not a and b['action'] == 'publish' and dist((a['lng'], a['lat']), (b['lng'], b['lat'])) < 3 for b in rows)

    # Toạ độ GPS do người khảo sát nhập tay: chỉ coi là "đã đối chiếu" khi có file GPX/KML khớp (--gps-proof=gpx)
    proof = next((a.split('=', 1)[1] for a in sys.argv if a.startswith('--gps-proof=')), 'pending')
    json.dump({'source': done_path.replace('\\', '/').split('/')[-1], 'convertedAt': datetime.now().isoformat(timespec='seconds'),
               'gpsProof': proof, 'rows': rows}, open(out_path, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    from collections import Counter
    print(f'{len(rows)} dòng -> {out_path}; hành động:', dict(Counter(r["action"] for r in rows)),
          '| GPS độc lập:', sum(r['gpsIndependent'] for r in rows), '| trùng điểm:', sum(r['sharedPoint'] for r in rows),
          '| lỗi:', problems)


if __name__ == '__main__':
    main(*[a for a in sys.argv[1:] if not a.startswith('--')][:3])
